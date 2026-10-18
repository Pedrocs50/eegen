# -*- coding: utf-8 -*-
"""Gerador de exercícios de listening: roteiros .txt -> áudios MP3 + planilha.

Instalação:
    pip install -r requirements.txt

Requer o FFmpeg no PATH (ou o ffmpeg.exe nesta mesma pasta).

Uso:
    python gerador.py                  # valida, gera áudios e planilha
    python gerador.py --validar        # só confere os roteiros
    python gerador.py --so-planilha    # não chama o TTS
    python gerador.py --so-audio --forcar --rate=-10% --pausa 600

Formato do roteiro (UTF-8; também aceita Windows-1252/Latin-1):

    TITLE: Greetings at the Cafe          (opcional; padrão: nome do arquivo)
    LEVEL: Below Beginner (Pre-A1)        (opcional; padrão: pelo código A0, A1...)
    VOICE MIA: en-GB-SoniaNeural          (opcional; escolhe a voz de um falante)

    AUDIO SCRIPT
    MIA: Hello.
    SAM: Hello.

    TRUE OR FALSE
    [T] Sam orders a small coffee.
    [F] Sam does not want milk.
    > Explicação opcional da questão acima.

    MULTIPLE CHOICE
    ? What does Sam order?
    * Coffee          (alternativa correta)
    - Tea
    - Juice

Linhas antes de ``AUDIO SCRIPT`` são metadados. No roteiro, cada fala fica em
``FALANTE: texto`` e linhas sem falante continuam a fala anterior. As
alternativas são embaralhadas de forma reproduzível (mude com ``--seed``).
"""

from __future__ import annotations

import argparse
import asyncio
import hashlib
import itertools
import json
import random
import re
import shutil
import subprocess
import tempfile
from dataclasses import dataclass, field
from pathlib import Path

BASE_DIR = Path(__file__).resolve().parent
PASTA_ROTEIROS = BASE_DIR / "roteiros"
PASTA_AUDIOS = BASE_DIR / "audios"
PLANILHA_PADRAO = BASE_DIR / "planilhas" / "exercicios.xlsx"
CACHE_AUDIOS = PASTA_AUDIOS / ".cache.json"
VOZES = [
    "en-GB-RyanNeural",      # britânico masculino
    "en-GB-SoniaNeural",     # britânica feminina
    "en-US-GuyNeural",       # americano masculino
    "en-US-JennyNeural",     # americana feminina
    "en-AU-NatashaNeural",   # australiana feminina
]
NIVEIS = {
    "A0": "Below Beginner (Pre-A1)",
    "A1": "Beginner (A1)",
    "A2": "Elementary (A2)",
    "B1": "Pre-Intermediate (B1)",
    "B2": "Intermediate (B2)",
    "C1": "Advanced (C1)",
    "C2": "Proficient (C2)",
}
LETRAS = "ABCDE"
TENTATIVAS_TTS = 3
SILENCIO_INICIAL_MS = 700
SECOES = {"AUDIO SCRIPT": "roteiro", "TRUE OR FALSE": "vf", "MULTIPLE CHOICE": "multipla"}
PADRAO_FALANTE = re.compile(r"^([^:]{1,40}):\s*(.+)$")
PADRAO_META = re.compile(r"^([A-Za-z][A-Za-z ]{0,40}[A-Za-z0-9]):\s*(.+)$")
PADRAO_VF = re.compile(r"^\[([TFVtfv])\]\s*(.+)$")
PADRAO_CODIGO = re.compile(r"^([A-C][0-2])_(\d+)_(.+)$")
edge_tts = None
FFMPEG: str | None = None


@dataclass
class Questao:
    enunciado: str
    alternativas: list[str]      # para V/F: ["True", "False"]
    correta: int                 # índice em ``alternativas``
    explicacao: str = ""


@dataclass
class Roteiro:
    caminho: Path
    codigo: str
    titulo: str
    nivel: str
    vozes: dict[str, str] = field(default_factory=dict)
    falas: list[tuple[str, str]] = field(default_factory=list)
    vf: list[Questao] = field(default_factory=list)
    multipla: list[Questao] = field(default_factory=list)
    erros: list[str] = field(default_factory=list)

    @property
    def nome(self) -> str:
        return str(self.caminho.relative_to(PASTA_ROTEIROS))

    @property
    def audio(self) -> Path:
        """Mantém subpastas de roteiros para evitar nomes de áudio duplicados."""
        return PASTA_AUDIOS / self.caminho.relative_to(PASTA_ROTEIROS).with_suffix(".mp3")


# --------------------------------------------------------------------------
# Leitura dos roteiros
# --------------------------------------------------------------------------

def ler_texto(caminho: Path) -> str:
    """Lê o arquivo tentando codificações comuns de roteiros no Windows."""
    dados = caminho.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return dados.decode(encoding)
        except UnicodeDecodeError:
            pass
    raise UnicodeError("Não foi possível identificar a codificação do arquivo.")


def metadados_padrao(caminho: Path) -> tuple[str, str, str]:
    """Deduz código, título e nível de nomes como ``A1_02_Bus_to_the_Park``."""
    encontrado = PADRAO_CODIGO.match(caminho.stem)
    if not encontrado:
        return caminho.stem, caminho.stem.replace("_", " "), ""
    prefixo, numero, resto = encontrado.groups()
    return f"{prefixo}_{numero}", resto.replace("_", " "), NIVEIS.get(prefixo, "")


def ler_roteiro(caminho: Path) -> Roteiro:
    """Lê metadados, falas e questões. Problemas vão para ``roteiro.erros``."""
    codigo, titulo, nivel = metadados_padrao(caminho)
    roteiro = Roteiro(caminho, codigo, titulo, nivel)
    linhas = ler_texto(caminho).splitlines()
    # Sem cabeçalho AUDIO SCRIPT, o arquivo já começa pelas falas (formato antigo).
    tem_cabecalho = any(linha.strip().upper().rstrip(":") == "AUDIO SCRIPT" for linha in linhas)
    secao = "meta" if tem_cabecalho else "roteiro"
    questao_aberta: Questao | None = None

    def erro(numero: int, mensagem: str) -> None:
        roteiro.erros.append(f"linha {numero}: {mensagem}")

    for numero, linha_original in enumerate(linhas, start=1):
        linha = linha_original.strip()
        if not linha:
            continue
        cabecalho = linha.upper().rstrip(":")
        if cabecalho in SECOES:
            secao = SECOES[cabecalho]
            questao_aberta = None
            continue

        if linha.startswith(">"):
            if secao in ("vf", "multipla") and questao_aberta:
                questao_aberta.explicacao = linha[1:].strip()
            else:
                erro(numero, "explicação '>' sem questão antes dela")
            continue

        if secao == "meta":
            encontrado = PADRAO_META.match(linha)
            if not encontrado:
                erro(numero, "metadado inválido (esperado CHAVE: valor)")
                continue
            chave, valor = encontrado.group(1).upper(), encontrado.group(2).strip()
            if chave == "TITLE":
                roteiro.titulo = valor
            elif chave == "LEVEL":
                roteiro.nivel = valor
            elif chave.startswith("VOICE "):
                roteiro.vozes[chave[6:].strip()] = valor
            else:
                erro(numero, f"metadado desconhecido: {chave}")

        elif secao == "roteiro":
            encontrado = PADRAO_FALANTE.match(linha)
            if encontrado:
                roteiro.falas.append((encontrado.group(1).strip(), encontrado.group(2).strip()))
            elif roteiro.falas:
                # Permite quebra de linha no meio de uma fala, sem descartar texto.
                falante, texto = roteiro.falas[-1]
                roteiro.falas[-1] = (falante, f"{texto} {linha}")
            else:
                erro(numero, "fala inválida (esperado FALANTE: texto)")

        elif secao == "vf":
            encontrado = PADRAO_VF.match(linha)
            if not encontrado:
                erro(numero, "verdadeiro/falso inválido (esperado [T] ou [F] antes da frase)")
                continue
            verdadeira = encontrado.group(1).upper() in "TV"
            questao_aberta = Questao(encontrado.group(2).strip(), ["True", "False"], 0 if verdadeira else 1)
            roteiro.vf.append(questao_aberta)

        elif secao == "multipla":
            marcador, texto = linha[0], linha[1:].strip()
            if marcador == "?" and texto:
                questao_aberta = Questao(texto, [], -1)
                roteiro.multipla.append(questao_aberta)
            elif marcador in "*-" and texto and questao_aberta:
                if marcador == "*":
                    if questao_aberta.correta >= 0:
                        erro(numero, "mais de uma alternativa correta (*)")
                    questao_aberta.correta = len(questao_aberta.alternativas)
                questao_aberta.alternativas.append(texto)
            else:
                erro(numero, "múltipla escolha inválida (use '? pergunta', '* correta', '- errada')")

    validar(roteiro)
    return roteiro


def validar(roteiro: Roteiro) -> None:
    if not roteiro.falas:
        roteiro.erros.append("nenhuma fala encontrada")
    if not roteiro.nivel:
        roteiro.erros.append("nível não identificado; adicione 'LEVEL: ...' no topo")
    if not roteiro.vf and not roteiro.multipla:
        roteiro.erros.append("nenhuma questão (seções TRUE OR FALSE / MULTIPLE CHOICE)")
    for numero, questao in enumerate(roteiro.multipla, start=1):
        quantidade = len(questao.alternativas)
        if not 2 <= quantidade <= len(LETRAS):
            roteiro.erros.append(f"múltipla escolha {numero}: precisa de 2 a {len(LETRAS)} alternativas, tem {quantidade}")
        if questao.correta < 0:
            roteiro.erros.append(f"múltipla escolha {numero}: marque a correta com '*'")
    falantes = {falante for falante, _ in roteiro.falas}
    for falante in roteiro.vozes.keys() - falantes:
        roteiro.erros.append(f"VOICE para falante inexistente: {falante}")


def embaralhar(roteiro: Roteiro, semente: str) -> None:
    """Embaralha as alternativas; a mesma semente gera sempre o mesmo gabarito."""
    sorteio = random.Random(f"{semente}:{roteiro.codigo}")
    for questao in roteiro.multipla:
        correta = questao.alternativas[questao.correta]
        sorteio.shuffle(questao.alternativas)
        questao.correta = questao.alternativas.index(correta)


# --------------------------------------------------------------------------
# Áudio
# --------------------------------------------------------------------------

def mapa_vozes(roteiro: Roteiro) -> dict[str, str]:
    falantes = list(dict.fromkeys(falante for falante, _ in roteiro.falas))
    livres = [voz for voz in VOZES if voz not in roteiro.vozes.values()] or VOZES
    automaticas = itertools.cycle(livres)
    return {falante: roteiro.vozes.get(falante) or next(automaticas) for falante in falantes}


def assinatura_audio(roteiro: Roteiro, pausa_ms: int, rate: str) -> str:
    """Muda só quando algo que afeta o áudio muda (editar questões não regera)."""
    dados = json.dumps([roteiro.falas, mapa_vozes(roteiro), pausa_ms, rate], ensure_ascii=False)
    return hashlib.sha256(dados.encode("utf-8")).hexdigest()


async def sintetizar(texto: str, voz: str, saida: Path, rate: str) -> None:
    """Gera um trecho e repete tentativas em falhas transitórias de rede."""
    ultimo_erro: Exception | None = None
    for tentativa in range(1, TENTATIVAS_TTS + 1):
        try:
            await edge_tts.Communicate(texto, voz, rate=rate).save(str(saida))
            if saida.exists() and saida.stat().st_size > 0:
                return
            raise RuntimeError("O serviço retornou um arquivo de áudio vazio.")
        except Exception as erro:  # edge-tts não expõe uma exceção única para rede.
            ultimo_erro = erro
            saida.unlink(missing_ok=True)
            if tentativa < TENTATIVAS_TTS:
                espera = tentativa * 2
                print(f"    tentativa {tentativa} falhou; repetindo em {espera}s...")
                await asyncio.sleep(espera)
    raise RuntimeError(f"TTS falhou após {TENTATIVAS_TTS} tentativas: {ultimo_erro}") from ultimo_erro


def juntar_trechos(trechos: list[Path], destino: Path, pausa_ms: int) -> None:
    """Concatena os trechos com o FFmpeg, com pausa entre falas e no início."""
    entradas = [argumento for trecho in trechos for argumento in ("-i", str(trecho))]
    pausas = ";".join(f"[{i}:a]apad=pad_dur={pausa_ms / 1000}[a{i}]" for i in range(len(trechos)))
    rotulos = "".join(f"[a{i}]" for i in range(len(trechos)))
    filtro = f"{pausas};{rotulos}concat=n={len(trechos)}:v=0:a=1,adelay={SILENCIO_INICIAL_MS}:all=1[saida]"
    comando = [FFMPEG, "-hide_banner", "-loglevel", "error", "-y", *entradas,
               "-filter_complex", filtro, "-map", "[saida]", "-c:a", "libmp3lame", "-q:a", "4", str(destino)]
    resultado = subprocess.run(comando, capture_output=True, text=True)
    if resultado.returncode != 0:
        raise RuntimeError(resultado.stderr.strip() or "FFmpeg falhou ao juntar os trechos.")


async def gerar_audio(roteiro: Roteiro, pausa_ms: int, rate: str) -> bool:
    """Gera um único MP3. Retorna False sem derrubar os outros roteiros."""
    vozes = mapa_vozes(roteiro)
    print(f"[+] {roteiro.nome}: {len(roteiro.falas)} falas, {len(vozes)} voz(es)")
    roteiro.audio.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="tts_", dir=BASE_DIR) as pasta_temporaria:
        trechos: list[Path] = []
        for indice, (falante, texto) in enumerate(roteiro.falas, start=1):
            trecho = Path(pasta_temporaria) / f"{indice:03d}.mp3"
            try:
                await sintetizar(texto, vozes[falante], trecho, rate)
                trechos.append(trecho)
            except Exception as erro:
                # Uma fala defeituosa não impede o restante do lote.
                print(f"    [!] fala {indice} de {falante} ignorada: {erro}")

        if not trechos:
            print(f"[ERRO] {roteiro.nome}: nenhuma fala pôde ser sintetizada.")
            return False

        arquivo_parcial = roteiro.audio.with_suffix(".part.mp3")
        try:
            juntar_trechos(trechos, arquivo_parcial, pausa_ms)
            arquivo_parcial.replace(roteiro.audio)
        except Exception as erro:
            arquivo_parcial.unlink(missing_ok=True)
            print(f"[ERRO] {roteiro.nome}: não foi possível salvar o MP3: {erro}")
            return False

    completo = len(trechos) == len(roteiro.falas)
    print(f"    -> {roteiro.audio.relative_to(BASE_DIR)} ({'ok' if completo else 'com falas ignoradas'})")
    # Áudio com falas faltando não entra no cache, para ser refeito na próxima vez.
    return completo


def ler_cache() -> dict[str, str]:
    try:
        return json.loads(CACHE_AUDIOS.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}


async def gerar_audios(roteiros: list[Roteiro], pausa_ms: int, rate: str, forcar: bool) -> int:
    """Gera os MP3s que mudaram desde a última vez. Retorna quantos falharam."""
    cache = ler_cache()
    falhas = 0
    for roteiro in roteiros:
        assinatura = assinatura_audio(roteiro, pausa_ms, rate)
        if not forcar and roteiro.audio.exists() and cache.get(roteiro.nome) == assinatura:
            print(f"[=] {roteiro.nome}: áudio já está atualizado")
            continue
        if await gerar_audio(roteiro, pausa_ms, rate):
            cache[roteiro.nome] = assinatura
            PASTA_AUDIOS.mkdir(parents=True, exist_ok=True)
            CACHE_AUDIOS.write_text(json.dumps(cache, indent=1, ensure_ascii=False), encoding="utf-8")
        else:
            falhas += 1
    return falhas


# --------------------------------------------------------------------------
# Planilha
# --------------------------------------------------------------------------

CABECALHO = ["Tipo", "Conteúdo", "Alt A", "Alt B", "Alt C", "Alt D", "Alt E",
             "Gabarito", "Explicação", "Link do Formulário:", "Link"]
LARGURAS = {"A": 15, "B": 48, "C": 20, "D": 20, "E": 20, "F": 20, "G": 20, "H": 13, "I": 36, "J": 20, "K": 20}


def explicacao_padrao(questao: Questao, verdadeiro_falso: bool) -> str:
    if questao.explicacao:
        return questao.explicacao
    resposta = questao.alternativas[questao.correta]
    if verdadeiro_falso:
        return f"The audio shows that this statement is {resposta.lower()}."
    return f"Correct answer: {resposta}."


def gerar_planilha(roteiros: list[Roteiro], destino: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    fonte = Font(name="Arial", size=10, color="1F1F1F")
    fonte_cabecalho = Font(name="Arial", bold=True, color="FFFFFF")
    fundo_cabecalho = PatternFill("solid", fgColor="1F4E78")
    quebra = Alignment(wrap_text=True, vertical="top")

    def estilizar_cabecalho(aba, cor_aba: str) -> None:
        aba.sheet_view.showGridLines = False
        aba.freeze_panes = "A2"
        aba.sheet_properties.tabColor = cor_aba
        for celula in aba[1]:
            celula.font = fonte_cabecalho
            celula.fill = fundo_cabecalho
            celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    livro = Workbook()
    indice = livro.active
    indice.title = "Roteiros"
    indice.append(["Código", "Nível", "Título", "Arquivo de áudio", "Falas", "V/F", "Múltipla escolha"])
    for roteiro in roteiros:
        indice.append([roteiro.codigo, roteiro.nivel, roteiro.titulo,
                       roteiro.audio.relative_to(BASE_DIR).as_posix(),
                       len(roteiro.falas), len(roteiro.vf), len(roteiro.multipla)])
    estilizar_cabecalho(indice, "1F4E78")
    for linha in indice.iter_rows(min_row=2):
        for celula in linha:
            celula.font, celula.alignment = fonte, quebra
    for coluna, largura in zip("ABCDEFG", (12, 26, 38, 46, 8, 8, 16)):
        indice.column_dimensions[coluna].width = largura

    numero_aba = 0
    for roteiro in roteiros:
        for sufixo, questoes, instrucao, verdadeiro_falso in (
            ("TF", roteiro.vf, "choose True or False", True),
            ("MC", roteiro.multipla, "choose the best answer", False),
        ):
            if not questoes:
                continue
            numero_aba += 1
            aba = livro.create_sheet(f"{numero_aba:02d}_{roteiro.codigo}_{sufixo}"[:31])
            aba.append(CABECALHO)
            audio = roteiro.audio.relative_to(BASE_DIR).as_posix()
            aba.append(["texto", f"Listening: {roteiro.titulo}",
                        f"Level: {roteiro.nivel}. Audio file: {audio}\nListen to the audio, then {instrucao}."])
            for numero, questao in enumerate(questoes, start=1):
                alternativas = questao.alternativas + [""] * (len(LETRAS) - len(questao.alternativas))
                aba.append(["pergunta", f"{numero}. {questao.enunciado}", *alternativas,
                            LETRAS[questao.correta], explicacao_padrao(questao, verdadeiro_falso)])

            estilizar_cabecalho(aba, "5B9BD5")
            for celula in aba[2]:
                celula.font = Font(name="Arial", italic=True, color="1F1F1F")
                celula.fill = PatternFill("solid", fgColor="EAF2F8")
                celula.alignment = quebra
            for linha in aba.iter_rows(min_row=3):
                for celula in linha:
                    celula.font, celula.alignment = fonte, quebra
            for coluna, largura in LARGURAS.items():
                aba.column_dimensions[coluna].width = largura
            aba.row_dimensions[2].height = 55

    destino.parent.mkdir(parents=True, exist_ok=True)
    livro.save(destino)
    print(f"Planilha: {destino.relative_to(BASE_DIR) if destino.is_relative_to(BASE_DIR) else destino} "
          f"({numero_aba} abas de exercícios)")


# --------------------------------------------------------------------------
# Principal
# --------------------------------------------------------------------------

def verificar_dependencias(precisa_audio: bool, precisa_planilha: bool) -> bool:
    """Carrega dependências e informa exatamente o que falta."""
    global edge_tts, FFMPEG
    try:
        if precisa_planilha:
            import openpyxl  # noqa: F401
        if precisa_audio:
            import edge_tts as edge_tts_modulo
            edge_tts = edge_tts_modulo
    except ModuleNotFoundError as erro:
        print(f"[ERRO] Pacote Python ausente: {erro.name}")
        print("Instale as dependências com: pip install -r requirements.txt")
        return False

    if not precisa_audio:
        return True
    local = BASE_DIR / "ffmpeg.exe"
    FFMPEG = shutil.which("ffmpeg") or (str(local) if local.exists() else None)
    if FFMPEG:
        return True
    print("[ERRO] FFmpeg não foi encontrado no PATH nem nesta pasta.")
    print("Instale o FFmpeg, feche e abra o terminal novamente, e execute o comando outra vez.")
    return False


async def main() -> int:
    parser = argparse.ArgumentParser(description="Gera áudios e planilha de exercícios a partir dos roteiros.")
    modo = parser.add_mutually_exclusive_group()
    modo.add_argument("--validar", action="store_true", help="só confere os roteiros, sem gerar nada")
    modo.add_argument("--so-audio", action="store_true", help="gera só os MP3s")
    modo.add_argument("--so-planilha", action="store_true", help="gera só a planilha")
    parser.add_argument("--forcar", action="store_true", help="regera todos os áudios, mesmo os atualizados")
    parser.add_argument("--pausa", type=int, default=450, help="silêncio entre falas, em ms (padrão: 450)")
    parser.add_argument("--rate", default="-5%", help="velocidade de fala do edge-tts (padrão: -5%%)")
    parser.add_argument("--seed", default="1", help="semente do embaralhamento das alternativas (padrão: 1)")
    parser.add_argument("--saida", type=Path, default=PLANILHA_PADRAO, help="caminho do .xlsx gerado")
    args = parser.parse_args()

    if args.pausa < 0:
        parser.error("--pausa deve ser zero ou maior.")
    if not PASTA_ROTEIROS.is_dir():
        print(f"[ERRO] Pasta de roteiros não encontrada: {PASTA_ROTEIROS}")
        return 1

    arquivos = sorted(PASTA_ROTEIROS.rglob("*.txt"), key=lambda caminho: str(caminho).lower())
    if not arquivos:
        print(f"[!] Nenhum arquivo .txt encontrado em: {PASTA_ROTEIROS}")
        return 0

    roteiros: list[Roteiro] = []
    com_erro = 0
    for arquivo in arquivos:
        try:
            roteiro = ler_roteiro(arquivo)
        except (OSError, UnicodeError) as erro:
            print(f"[ERRO] {arquivo.name}: não foi possível ler o roteiro: {erro}")
            com_erro += 1
            continue
        if roteiro.erros:
            com_erro += 1
            print(f"[ERRO] {roteiro.nome}:")
            for mensagem in roteiro.erros:
                print(f"    - {mensagem}")
        roteiros.append(roteiro)

    print(f"Encontrados {len(arquivos)} roteiro(s); {len(arquivos) - com_erro} sem erros.")
    if com_erro:
        print("Corrija os erros acima antes de gerar os materiais.")
        return 1
    if args.validar:
        return 0

    precisa_audio = not args.so_planilha
    precisa_planilha = not args.so_audio
    if not verificar_dependencias(precisa_audio, precisa_planilha):
        return 1

    falhas = 0
    if precisa_audio:
        falhas = await gerar_audios(roteiros, args.pausa, args.rate, args.forcar)
    if precisa_planilha:
        for roteiro in roteiros:
            embaralhar(roteiro, args.seed)
        gerar_planilha(roteiros, args.saida.resolve())

    if falhas:
        print(f"\nConcluído com {falhas} áudio(s) com problema.")
        return 2
    print("\nConcluído.")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(asyncio.run(main()))
    except KeyboardInterrupt:
        print("\nGeração cancelada pelo usuário.")
        raise SystemExit(130)
